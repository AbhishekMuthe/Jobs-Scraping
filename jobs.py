import scrapy
import re
from scrapy.crawler import CrawlerProcess
from openpyxl.workbook import Workbook
import requests
import pandas as pd

class Spider(scrapy.Spider):
    name = 'internships'
    allowed_domains = ['naukri.com']
    start_urls = ['https://www.naukri.com/']
    def parse(self, response):
        url = "https://www.naukri.com/information-technology-jobs"
        yield scrapy.Request(url, callback=self.parse2)

    def parse2(self, response):
        count = 0
        title = []
        company_name = []
        company_id = []
        description = []
        years = []
        salary = []
        headers = {'appid': "109",
           'systemid': "109"}
        for i in range(1,100):
            url1 = "https://www.naukri.com/jobapi/v3/search?noOfResults=&urlType=search_by_keyword&searchType=adv&keyword=information%20technology&pageNo=" + str(i) + "&seoKey=information-technology-jobs&src=jobsearchDesk&latLong="
            try:
                req = requests.get(url1, headers=headers)
                data = req.json()
            except:
                break
            data = data['jobDetails']
            for details in data:
                title.append(details['title'])
                company_name.append(details['companyName'])
                desc = cleanhtml(details['jobDescription'])
                description.append(desc)
                company_id.append(details['companyId'])
                # for placeholder in details['placeholders']:
                experience = details['placeholders'][0]['label']
                years.append(experience)
                amount = details['placeholders'][1]['label']
                salary.append(amount)
                count = count + 1
                if(count==100):
                    break
            if(count==100):
                break
        df = pd.DataFrame({"Company Name": company_name, "Title": title, "Description": description, "Company Id": company_id, "Salary": salary, "Experience": years})
        df.to_excel('jobs.xlsx', index=False)
        
def cleanhtml(raw_html):
    cleanr = re.compile('<.*?>')
    cleantext = re.sub(cleanr, '', raw_html)
    return cleantext



crawlprocess = CrawlerProcess({
    'USER_AGENT': 'Mozilla/5.0 (Windows NT 6.1; WOW64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/34.0.1847.131 Safari/537.36',
    'FEED_FORMAT': 'csv',
    'FEED_URI': 'output.csv',
    'DEPTH_LIMIT': 1000000,
    'CLOSESPIDER_PAGECOUNT': 10000000,
})
crawlprocess.crawl(Spider, urls_file='input.txt')
crawlprocess.start()